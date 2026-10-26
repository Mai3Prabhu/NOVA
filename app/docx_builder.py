"""Builds the final Word (.docx) document from the agent's results."""

import re
import uuid
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph

GENERATED_DIR = Path(__file__).resolve().parent.parent / "generated"

# Recognized Markdown patterns only - never a blanket asterisk strip. A lone
# "*" or an unclosed "**" is left in the text exactly as written.
BOLD_PATTERN = re.compile(r"\*\*(.+?)\*\*")
BULLET_LINE_PATTERN = re.compile(r"^[-*]\s+(.*)$")
NUMBERED_LINE_PATTERN = re.compile(r"^\d+[.)]\s+(.*)$")


def _apply_base_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for section in document.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def _add_inline_runs(paragraph: Paragraph, text: str, *, force_bold: bool = False) -> None:
    """Add `text` to `paragraph`, turning **bold** markers into bold runs.

    Everything outside a matched `**...**` pair is added as plain text, so
    ordinary asterisks (a single "*", or an unclosed "**") pass through
    unchanged instead of being stripped.
    """
    pos = 0
    matched = False
    for match in BOLD_PATTERN.finditer(text):
        matched = True
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()]).bold = force_bold
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    remainder = text[pos:]
    if remainder or not matched:
        paragraph.add_run(remainder).bold = force_bold


def _set_cell_text(cell: _Cell, text: str, *, bold: bool = False) -> None:
    _add_inline_runs(cell.paragraphs[0], text, force_bold=bold)


def _add_markdown_block(document: Document, text: str) -> None:
    """Render multi-line text, converting Markdown bullet/numbered lines into
    real Word list paragraphs, and **bold** markers into bold runs.
    """
    any_line = False
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        any_line = True

        bullet_match = BULLET_LINE_PATTERN.match(stripped)
        numbered_match = NUMBERED_LINE_PATTERN.match(stripped)
        if bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, bullet_match.group(1))
        elif numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline_runs(paragraph, numbered_match.group(1))
        else:
            paragraph = document.add_paragraph()
            _add_inline_runs(paragraph, stripped)

    if not any_line:
        document.add_paragraph(text)


def _add_bulleted_list(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        _add_inline_runs(paragraph, item)


def _add_sections(document: Document, sections: list[dict]) -> None:
    document.add_heading("Document Content", level=1)
    if not sections:
        document.add_paragraph("No sections were generated for this document.")
        return
    for section in sections:
        heading = str(section.get("section") or "Section").replace("_", " ").title()
        content = section.get("content") or ""
        document.add_heading(heading, level=2)
        _add_markdown_block(document, content if content else "(no content generated)")


def _add_tables(document: Document, tables: list[dict]) -> None:
    if not tables:
        return
    document.add_heading("Tables", level=1)
    for table in tables:
        name = str(table.get("table_name") or "Table").replace("_", " ").title()
        headers = table.get("headers") or []
        rows = table.get("rows") or []
        document.add_heading(name, level=2)

        if not headers:
            document.add_paragraph("(no table data available)")
            continue

        word_table = document.add_table(rows=1, cols=len(headers))
        word_table.style = "Light Grid Accent 1"
        for cell, header in zip(word_table.rows[0].cells, headers):
            _set_cell_text(cell, str(header), bold=True)
        for row in rows:
            row_cells = word_table.add_row().cells
            for cell, value in zip(row_cells, row):
                _set_cell_text(cell, str(value))


def build_docx(
    doc_type: str,
    original_request: str,
    assumptions: list[str],
    open_questions: list[str],
    sections: list[dict],
    tables: list[dict],
) -> tuple[str, str]:
    """Render the plan's results into a .docx file under GENERATED_DIR.

    Returns (document_id, filename). The filename is always derived from a
    freshly generated UUID, never from user input, so it can't be used to
    traverse or overwrite arbitrary paths.
    """
    GENERATED_DIR.mkdir(parents=True, exist_ok=True)

    document = Document()
    _apply_base_styles(document)

    title = (doc_type or "document").replace("_", " ").strip().title() or "Generated Document"
    document.add_heading(title, level=0)

    document.add_heading("Overview", level=1)
    overview_paragraph = document.add_paragraph()
    _add_inline_runs(
        overview_paragraph,
        f'This {title.lower()} was generated in response to the following request: '
        f'"{original_request}"',
    )

    if assumptions:
        document.add_heading("Assumptions", level=1)
        document.add_paragraph(
            "The request did not specify these details. The items below are "
            "illustrative assumptions, not verified facts, and should be "
            "confirmed before this document is relied upon."
        )
        _add_bulleted_list(document, assumptions)

    _add_sections(document, sections)
    _add_tables(document, tables)

    if open_questions:
        document.add_heading("Open Questions", level=1)
        document.add_paragraph(
            "These details could not be reasonably assumed and should be "
            "clarified with the requester before finalizing this document."
        )
        _add_bulleted_list(document, open_questions)

    document_id = uuid.uuid4().hex
    filename = f"{document_id}.docx"
    document.save(str(GENERATED_DIR / filename))

    return document_id, filename
