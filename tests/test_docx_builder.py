from docx import Document

from app import docx_builder


def _all_text(document: Document) -> str:
    paragraph_text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    return paragraph_text + "\n" + table_text


def test_build_docx_creates_valid_docx_with_expected_content(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    document_id, filename = docx_builder.build_docx(
        doc_type="project_proposal",
        original_request="Create a project proposal for an AI chatbot.",
        assumptions=["Audience is internal stakeholders."],
        open_questions=["What is the budget?"],
        sections=[{"section": "overview", "content": "This project improves support."}],
        tables=[
            {
                "table_name": "timeline",
                "headers": ["Milestone", "Date"],
                "rows": [["Kickoff", "Jan 1"], ["Launch", "Mar 1"]],
            }
        ],
    )

    file_path = tmp_path / filename
    assert file_path.is_file()
    assert filename == f"{document_id}.docx"

    document = Document(str(file_path))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    text = _all_text(document)

    assert any("Project Proposal" in h for h in headings)
    assert "Create a project proposal for an AI chatbot." in text
    assert "This project improves support." in text
    assert "Audience is internal stakeholders." in text
    assert "What is the budget?" in text

    assert len(document.tables) == 1
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == ["Milestone", "Date"]
    assert [cell.text for cell in table.rows[1].cells] == ["Kickoff", "Jan 1"]


def test_build_docx_handles_missing_sections_and_tables_gracefully(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    document_id, filename = docx_builder.build_docx(
        doc_type="report",
        original_request="Summarize last quarter.",
        assumptions=[],
        open_questions=[],
        sections=[],
        tables=[],
    )

    file_path = tmp_path / filename
    assert file_path.is_file()
    document = Document(str(file_path))
    assert len(document.tables) == 0  # no crash, no table rendered


def test_build_docx_creates_generated_dir_if_missing(monkeypatch, tmp_path):
    missing_dir = tmp_path / "does" / "not" / "exist"
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", missing_dir)

    document_id, filename = docx_builder.build_docx(
        doc_type="report",
        original_request="Summarize last quarter.",
        assumptions=[],
        open_questions=[],
        sections=[],
        tables=[],
    )

    assert (missing_dir / filename).is_file()


def test_build_docx_filenames_are_unique_and_not_derived_from_input(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    args = dict(
        doc_type="report",
        original_request="../../etc/passwd",
        assumptions=[],
        open_questions=[],
        sections=[],
        tables=[],
    )
    id_1, filename_1 = docx_builder.build_docx(**args)
    id_2, filename_2 = docx_builder.build_docx(**args)

    assert id_1 != id_2
    assert "/" not in filename_1 and ".." not in filename_1


def _runs(document: Document):
    return [run for paragraph in document.paragraphs for run in paragraph.runs]


def test_markdown_bold_renders_as_bold_runs_without_asterisks(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    _, filename = docx_builder.build_docx(
        doc_type="report",
        original_request="Explain the system.",
        assumptions=["**Audience** is internal stakeholders."],
        open_questions=[],
        sections=[{"section": "details", "content": "**Functional Boundaries** define scope."}],
        tables=[],
    )

    document = Document(str(tmp_path / filename))
    runs = _runs(document)

    bold_run = next(r for r in runs if r.text == "Functional Boundaries")
    assert bold_run.bold is True

    bold_assumption_run = next(r for r in runs if r.text == "Audience")
    assert bold_assumption_run.bold is True

    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "**" not in full_text
    assert "Functional Boundaries define scope." in full_text
    assert "Audience is internal stakeholders." in full_text


def test_markdown_bullet_and_numbered_lines_become_word_lists(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    content = (
        "Intro line before any list.\n"
        "- **Data Integration**: connects to the CRM\n"
        "- Reporting: weekly summaries\n"
        "1. First step is critical\n"
        "2. Second step\n"
    )

    _, filename = docx_builder.build_docx(
        doc_type="report",
        original_request="Explain the system.",
        assumptions=[],
        open_questions=[],
        sections=[{"section": "details", "content": content}],
        tables=[],
    )

    document = Document(str(tmp_path / filename))

    bullet_paragraphs = [p for p in document.paragraphs if p.style.name == "List Bullet"]
    numbered_paragraphs = [p for p in document.paragraphs if p.style.name == "List Number"]

    assert any(p.text == "Data Integration: connects to the CRM" for p in bullet_paragraphs)
    assert any(p.text == "Reporting: weekly summaries" for p in bullet_paragraphs)
    assert any(p.text == "First step is critical" for p in numbered_paragraphs)
    assert any(p.text == "Second step" for p in numbered_paragraphs)

    # The raw Markdown markers must not survive into any paragraph's text.
    for paragraph in document.paragraphs:
        assert not paragraph.text.lstrip().startswith("- ")
        assert not paragraph.text.lstrip().startswith("1. ")
        assert "**" not in paragraph.text


def test_ordinary_asterisks_are_left_untouched(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    content = (
        "Use caution: 3 * 4 = 12 is just math, not markdown.\n"
        "An unclosed **marker stays literal.\n"
        "What about *single-star emphasis* staying as-is?\n"
    )

    _, filename = docx_builder.build_docx(
        doc_type="report",
        original_request="Explain the system.",
        assumptions=[],
        open_questions=[],
        sections=[{"section": "details", "content": content}],
        tables=[],
    )

    document = Document(str(tmp_path / filename))
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert "3 * 4 = 12 is just math, not markdown." in full_text
    assert "An unclosed **marker stays literal." in full_text
    assert "What about *single-star emphasis* staying as-is?" in full_text


def test_existing_headings_tables_assumptions_and_open_questions_still_work(monkeypatch, tmp_path):
    monkeypatch.setattr(docx_builder, "GENERATED_DIR", tmp_path)

    document_id, filename = docx_builder.build_docx(
        doc_type="project_proposal",
        original_request="Create a project proposal for an AI chatbot.",
        assumptions=["**Audience** is internal stakeholders."],
        open_questions=["What is the **budget**?"],
        sections=[{"section": "overview", "content": "**Scope**: internal support only."}],
        tables=[
            {
                "table_name": "timeline",
                "headers": ["**Milestone**", "Date"],
                "rows": [["Kickoff", "Jan 1"]],
            }
        ],
    )

    document = Document(str(tmp_path / filename))
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert any("Project Proposal" in h for h in headings)
    assert any("Assumptions" == h for h in headings)
    assert any("Open Questions" == h for h in headings)
    assert "Audience is internal stakeholders." in full_text
    assert "What is the budget?" in full_text
    assert "Scope: internal support only." in full_text

    table = document.tables[0]
    header_cells = table.rows[0].cells
    assert header_cells[0].text == "Milestone"
    header_run = header_cells[0].paragraphs[0].runs[0]
    assert header_run.bold is True
