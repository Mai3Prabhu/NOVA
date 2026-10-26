"""Tests for the LLM-call/token reductions: combined section/table calls,
output-token caps, truncation handling and bounded correction.

All provider calls are mocked - no live Groq or Gemini access.
"""

import re
from typing import ClassVar
from types import SimpleNamespace

import pytest
from pydantic import BaseModel

from app.executor import execute_plan
from app.llm import LLMError, SchemaValidationError, request_structured_json
from app.models import AgentPlan
from app.tools import (
    IncompleteOutputError,
    SectionContent,
    SectionItem,
    SectionsBatch,
    TableContent,
    TableItem,
    TablesBatch,
    ToolError,
    write_section,
)

USER_REQUEST = "Create a project proposal for an AI customer support chatbot."


def _task(task_id: str, tool_name: str, **args) -> dict:
    return {"task_id": task_id, "description": f"do {task_id}", "tool_name": tool_name, "args": args}


def _complex_plan() -> AgentPlan:
    """The maximum-size plan: gather + 4 sections + 2 tables + build = 8."""
    return AgentPlan.model_validate(
        {
            "doc_type": "project_proposal",
            "assumptions": ["Audience is internal stakeholders (not specified by the user)."],
            "open_questions": ["What is the budget?"],
            "tasks": [
                _task("t1", "gather_context", topic="chatbot"),
                _task("t2", "write_section", section="overview"),
                _task("t3", "write_section", section="scope"),
                _task("t4", "write_section", section="timeline"),
                _task("t5", "write_section", section="risks"),
                _task("t6", "make_table", table_name="milestones"),
                _task("t7", "make_table", table_name="risk_register"),
                _task("t8", "build_docx"),
            ],
        }
    )


class FakeLLM:
    """Stands in for app.tools.request_structured_json and records each call."""

    def __init__(self, responder=None):
        self.calls = []
        self.responder = responder

    def __call__(self, system_prompt, user_prompt, response_model, max_tokens=None):
        self.calls.append(
            {"model": response_model, "user": user_prompt, "max_tokens": max_tokens}
        )
        if self.responder is not None:
            return self.responder(len(self.calls), user_prompt, response_model)
        return self.well_behaved(user_prompt, response_model)

    @staticmethod
    def well_behaved(user_prompt, response_model):
        ids = re.findall(r"task_id: (\S+)", user_prompt)
        if response_model is SectionsBatch:
            return SectionsBatch(
                sections=[SectionItem(task_id=i, content=f"text for {i}") for i in ids]
            )
        if response_model is TablesBatch:
            return TablesBatch(
                tables=[TableItem(task_id=i, headers=["A", "B"], rows=[[i, "x"]]) for i in ids]
            )
        if response_model is SectionContent:
            return SectionContent(content="single section")
        return TableContent(headers=["A"], rows=[["single"]])


# --- Combined generation -----------------------------------------------------


def test_complex_plan_uses_one_call_per_kind_not_one_per_task(monkeypatch, tmp_path):
    monkeypatch.setattr("app.docx_builder.GENERATED_DIR", tmp_path)
    llm = FakeLLM()
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    context = execute_plan(USER_REQUEST, _complex_plan())

    # 4 sections + 2 tables used to be 6 calls; now 2.
    assert sorted(c["model"].__name__ for c in llm.calls) == ["SectionsBatch", "TablesBatch"]


def test_combined_results_keep_plan_order_and_docx_builder_shape(monkeypatch, tmp_path):
    from docx import Document

    monkeypatch.setattr("app.docx_builder.GENERATED_DIR", tmp_path)
    monkeypatch.setattr("app.tools.request_structured_json", FakeLLM())

    plan = _complex_plan()
    context = execute_plan(USER_REQUEST, plan)

    assert [s["section"] for s in context["sections"]] == ["overview", "scope", "timeline", "risks"]
    assert [t["table_name"] for t in context["tables"]] == ["milestones", "risk_register"]
    # Exactly the keys docx_builder consumes, same as the single-task tools.
    assert set(context["sections"][0]) == {"section", "content"}
    assert set(context["tables"][0]) == {"table_name", "headers", "rows"}
    assert [e["status"] for e in context["execution_log"]] == ["success"] * 8

    filename = context["task_results"]["t8"]["filename"]
    document = Document(str(tmp_path / filename))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "What is the budget?" in text  # open questions survive
    assert len(document.tables) >= 2


def test_combined_prompt_keeps_grounding_rules_and_labels_once(monkeypatch, tmp_path):
    monkeypatch.setattr("app.docx_builder.GENERATED_DIR", tmp_path)
    prompts = []

    def spy(system_prompt, user_prompt, response_model, max_tokens=None):
        prompts.append((system_prompt, user_prompt))
        return FakeLLM.well_behaved(user_prompt, response_model)

    monkeypatch.setattr("app.tools.request_structured_json", spy)

    execute_plan(USER_REQUEST, _complex_plan())

    for system_prompt, user_prompt in prompts:
        assert "never state an invented number" in system_prompt  # grounding rules present
        assert user_prompt.count("Original user request") == 1  # sent once, not per section
        assert USER_REQUEST in user_prompt
        assert "What is the budget?" in user_prompt  # open questions reach the model
        assert "not supplied by the user" in user_prompt  # mock context stays labeled


def test_single_section_still_uses_the_single_tool(monkeypatch):
    llm = FakeLLM()
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    from app.models import PlanTask

    task = PlanTask(task_id="t1", description="d", tool_name="write_section", args={"section": "s"})
    result = write_section(task, {"request": USER_REQUEST, "assumptions": []})

    assert result == {"section": "s", "content": "single section"}
    assert [c["model"] for c in llm.calls] == [SectionContent]
    assert llm.calls[0]["max_tokens"] is None  # cap comes from the model class instead


# --- Bounded correction ------------------------------------------------------


def test_combined_response_missing_a_section_gets_exactly_one_correction(monkeypatch, tmp_path):
    monkeypatch.setattr("app.docx_builder.GENERATED_DIR", tmp_path)

    def responder(call_number, user_prompt, response_model):
        if response_model is SectionsBatch and call_number == 1:
            # Drops the last requested section.
            ids = re.findall(r"task_id: (\S+)", user_prompt)[:-1]
            return SectionsBatch(sections=[SectionItem(task_id=i, content="x") for i in ids])
        return FakeLLM.well_behaved(user_prompt, response_model)

    llm = FakeLLM(responder)
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    context = execute_plan(USER_REQUEST, _complex_plan())

    assert len(context["sections"]) == 4  # nothing silently dropped
    batch_calls = [c for c in llm.calls if c["model"] is SectionsBatch]
    assert len(batch_calls) == 2  # original + one correction
    assert "invalid for this reason" in batch_calls[1]["user"]


def test_correction_is_bounded_and_failure_is_recorded_not_hidden(monkeypatch, tmp_path):
    from app.executor import ExecutionError

    monkeypatch.setattr("app.docx_builder.GENERATED_DIR", tmp_path)

    def always_drops_one(call_number, user_prompt, response_model):
        if response_model is SectionsBatch:
            ids = re.findall(r"task_id: (\S+)", user_prompt)[:-1]
            return SectionsBatch(sections=[SectionItem(task_id=i, content="x") for i in ids])
        return FakeLLM.well_behaved(user_prompt, response_model)

    llm = FakeLLM(always_drops_one)
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    with pytest.raises(ExecutionError) as exc_info:
        execute_plan(USER_REQUEST, _complex_plan())

    assert len([c for c in llm.calls if c["model"] is SectionsBatch]) == 2  # no third attempt
    log = {e["task_id"]: e for e in exc_info.value.execution_log}
    assert all(log[t]["status"] == "failed" for t in ("t2", "t3", "t4", "t5"))
    assert log["t6"]["status"] == "success"  # tables call is independent
    assert "t8" not in log  # build_docx never ran


def test_single_tool_schema_error_gets_one_correction(monkeypatch):
    def responder(call_number, user_prompt, response_model):
        if call_number == 1:
            raise SchemaValidationError("rows must match headers", {"headers": ["A"], "rows": [["1", "2"]]})
        return TableContent(headers=["A"], rows=[["1"]])

    llm = FakeLLM(responder)
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    from app.models import PlanTask
    from app.tools import make_table

    task = PlanTask(task_id="t1", description="d", tool_name="make_table", args={"table_name": "t"})
    result = make_table(task, {"request": USER_REQUEST})

    assert result["rows"] == [["1"]]
    assert len(llm.calls) == 2


def test_non_schema_llm_error_is_not_corrected(monkeypatch):
    def responder(call_number, user_prompt, response_model):
        raise LLMError("both providers down")

    llm = FakeLLM(responder)
    monkeypatch.setattr("app.tools.request_structured_json", llm)

    from app.models import PlanTask

    task = PlanTask(task_id="t1", description="d", tool_name="write_section", args={})
    with pytest.raises(ToolError):
        write_section(task, {"request": USER_REQUEST})

    assert len(llm.calls) == 1  # no pointless second call


def test_incomplete_output_error_is_a_schema_error():
    assert issubclass(IncompleteOutputError, SchemaValidationError)


# --- Output-token caps and truncation ----------------------------------------


class Dummy(BaseModel):
    content: str


class CappedDummy(Dummy):
    max_output_tokens: ClassVar[int] = 123


def _groq_completion(content, finish_reason="stop"):
    return SimpleNamespace(
        choices=[
            SimpleNamespace(
                message=SimpleNamespace(content=content), finish_reason=finish_reason
            )
        ],
        usage=SimpleNamespace(prompt_tokens=10, completion_tokens=5),
    )


class RecordingGroq:
    def __init__(self, completions):
        self.completions = list(completions)
        self.kwargs = []
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kwargs):
        self.kwargs.append(kwargs)
        effect = self.completions.pop(0)
        if isinstance(effect, Exception):
            raise effect
        return effect


def _patch_groq(monkeypatch, completions):
    client = RecordingGroq(completions)
    monkeypatch.setattr("app.llm.get_groq_client", lambda: client)
    monkeypatch.setattr("app.llm.time.sleep", lambda s: None)
    return client


def test_model_declared_cap_is_sent_to_groq(monkeypatch):
    client = _patch_groq(monkeypatch, [_groq_completion('{"content": "x"}')])

    request_structured_json("s", "u", CappedDummy)

    assert client.kwargs[0]["max_completion_tokens"] == 123


def test_explicit_cap_overrides_model_cap_and_no_cap_sends_none(monkeypatch):
    client = _patch_groq(
        monkeypatch, [_groq_completion('{"content": "x"}'), _groq_completion('{"content": "x"}')]
    )

    request_structured_json("s", "u", CappedDummy, max_tokens=500)
    request_structured_json("s", "u", Dummy)

    assert client.kwargs[0]["max_completion_tokens"] == 500
    assert "max_completion_tokens" not in client.kwargs[1]


def test_real_models_declare_caps():
    assert AgentPlan.max_output_tokens > 0
    assert SectionContent.max_output_tokens > 0
    assert TableContent.max_output_tokens > 0


def test_truncated_groq_response_is_never_parsed_and_falls_back(monkeypatch):
    client = _patch_groq(monkeypatch, [_groq_completion('{"content": "cut o', "length")])
    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    monkeypatch.setattr("app.llm._call_gemini", lambda s, u: '{"content": "from gemini"}')

    result = request_structured_json("s", "u", Dummy)

    assert result.content == "from gemini"
    assert len(client.kwargs) == 1  # the same cap would truncate again: no Groq retry


def test_truncated_groq_response_without_fallback_fails_clearly(monkeypatch):
    _patch_groq(monkeypatch, [_groq_completion('{"content": "cut o', "length")])
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    with pytest.raises(LLMError, match="truncated"):
        request_structured_json("s", "u", Dummy)


def test_gemini_max_tokens_finish_reason_is_an_explicit_permanent_error(monkeypatch):
    from app.llm import _call_gemini, _PermanentProviderError

    class FakeModels:
        def generate_content(self, **kwargs):
            candidate = SimpleNamespace(finish_reason=SimpleNamespace(name="MAX_TOKENS"))
            return SimpleNamespace(candidates=[candidate], text='{"content": "cut')

    class FakeClient:
        def __init__(self, **kwargs):
            self.models = FakeModels()

    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    monkeypatch.setattr("google.genai.Client", FakeClient)

    with pytest.raises(_PermanentProviderError, match="truncated"):
        _call_gemini("s", "u", max_tokens=100)


# --- Rate limits stay bounded -------------------------------------------------


def test_repeated_rate_limits_make_a_bounded_number_of_calls(monkeypatch):
    import groq
    import httpx

    request = httpx.Request("POST", "https://api.groq.com/openai/v1/chat/completions")

    def limited():
        response = httpx.Response(429, request=request, headers={"retry-after": "2"})
        return groq.RateLimitError("rate limited", response=response, body=None)

    client = _patch_groq(monkeypatch, [limited(), limited(), limited(), limited()])
    sleeps = []
    monkeypatch.setattr("app.llm.time.sleep", lambda s: sleeps.append(s))
    monkeypatch.setenv("GEMINI_API_KEY", "fake")
    gemini_calls = []

    def gemini_limited(system_prompt, user_prompt):
        from app.llm import _PermanentProviderError

        gemini_calls.append(1)
        raise _PermanentProviderError("Gemini request failed (permanent): 429")

    monkeypatch.setattr("app.llm._call_gemini", gemini_limited)

    with pytest.raises(LLMError):
        request_structured_json("s", "u", Dummy)

    assert len(client.kwargs) == 2  # one Groq retry, then stop
    assert len(gemini_calls) == 1  # permanent Gemini error: no retry
    assert sleeps == [2.0]  # one bounded, provider-guided wait
